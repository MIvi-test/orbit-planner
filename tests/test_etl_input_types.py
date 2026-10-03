"""Excel не должен превращать повреждённые данные в нули и NULL."""
from decimal import Decimal

import openpyxl
import pytest

from etl import load


@pytest.fixture(autouse=True)
def reset_quality():
    load.dq.rows = []
    yield
    load.dq.rows = []


def sheet_with(value, *, formula=None):
    workbook = openpyxl.Workbook()
    ws = workbook.active
    ws.title = "Input"
    ws["B2"] = value
    formulas = openpyxl.Workbook().active
    if formula is not None:
        formulas["B2"] = formula
    return load.Sheet(ws, formulas)


@pytest.mark.parametrize(("value", "rule"), [
    (None, "MISSING_NUMBER"),
    (-1, "NEGATIVE_NUMBER"),
    (0.5, "FRACTIONAL_NUMBER"),
    ("NaN", "INVALID_NUMBER"),
])
def test_required_sp_reports_cell_instead_of_becoming_zero(value, rule):
    sheet = sheet_with(value)
    assert sheet.number(2, 2, required=True, nonnegative=True, integer=True) is None
    assert load.dq.rows[0][1] == "Input!B2"
    assert load.dq.rows[0][2] == rule


def test_formula_without_cached_result_is_error():
    sheet = sheet_with(None, formula="=1+2")
    assert sheet.number(2, 2, required=True) is None
    assert "формулы" in load.dq.rows[0][-1]


def test_negative_and_overprecise_hours_are_errors():
    for value, rule in [(-0.5, "NEGATIVE_NUMBER"), (0.001, "EXCESS_PRECISION")]:
        sheet = sheet_with(value)
        assert sheet.number(2, 2, nonnegative=True, decimal_places=2) is None
        assert load.dq.rows[-1][2] == rule
    assert sheet_with("1.000").number(2, 2, nonnegative=True, decimal_places=2) == Decimal(1)


def test_invalid_date_reports_cell():
    sheet = sheet_with("31.02.2026")
    assert sheet.date(2, 2) is None
    assert load.dq.rows[0][1:3] == ("Input!B2", "INVALID_DATE")


def test_duplicate_or_missing_column_has_locations():
    sheet = sheet_with(None)
    sheet.ws["A1"] = "task_id"
    sheet.ws["B1"] = "task_id"
    with pytest.raises(load.DataQualityError, match="блокирующие ошибки") as exc:
        sheet.columns(1)
    assert "A1" in exc.value.problems[0] and "B1" in exc.value.problems[0]

    sheet.ws["B1"] = "status"
    with pytest.raises(load.DataQualityError) as exc:
        sheet.require_columns(sheet.columns(1), 1, ("task_id", "estimation_sp"))
    assert "estimation_sp" in exc.value.problems[0]
